#!/usr/bin/env python
# -*- coding: utf-8 -*-
# @Date    : 2017-07-11 10:33:46
# @Author  : ditto (969956574@qq.com)
# @Link    : https://github.com/dittoyy
# @Version : $Id$

from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

wb = Workbook()
ws = wb.active

data = [
    ['Apples', 10000, 5000, 8000, 6000],
    ['Pears',   2000, 3000, 4000, 5000],
    ['Bananas', 6000, 6000, 6500, 6000],
    ['Oranges',  500,  300,  200,  700],
]

# add column headings. NB. these must be strings
ws.append(["Fruit", "2011", "2012", "2013", "2014"])
for row in data:
    ws.append(row)

tab = Table(displayName="Table1", ref="A1:E5")

# Add a default style with striped rows and banded columns
style = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False,
                       showLastColumn=False, showRowStripes=True, showColumnStripes=True)
tab.tableStyleInfo = style
ws.add_table(tab)
# wb.save("table.xlsx")

# #Accessing a range called “my_range”:
# my_range = wb.defined_names['my_range']
# # if this contains a range of cells then the destinations attribute is not None
# dests = my_range.destinations # returns a generator of (worksheet title, cell range) tuples

# cells = []
# for title, coord in dests:
#     ws = wb[title]
#     cells.append(ws[coord])
# wb.save("table.xlsx")

